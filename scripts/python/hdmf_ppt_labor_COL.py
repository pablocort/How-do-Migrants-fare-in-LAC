#!/usr/bin/env python3
"""
hdmf_ppt_labor_COL.py -- Build a PowerPoint presentation for COL labor trends.

One chart per slide, summary tables from Excel, timing table at the end.
Aesthetics follow harmonization_System/skills/presentation_aesthetics.md.

Usage:
    py hdmf_ppt_labor_COL.py
"""

print("File created with the Claude HDMF system -- 2026-04-08")

import os
import pandas as pd
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

# =============================================================================
# CONFIGURATION
# =============================================================================

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DATE_TAG  = '2026-04-10'
COUNTRY   = 'COL'

TREND_DIR  = os.path.join(BASE_DIR, 'out', 'trends', DATE_TAG, COUNTRY)
LABOR_DIR  = os.path.join(TREND_DIR, 'labor')
EXCEL_FILE = os.path.join(BASE_DIR, 'out', 'indicator_descriptive', DATE_TAG,
                          f'hdmf_trends_{COUNTRY}.xlsx')
OUT_DIR    = os.path.join(BASE_DIR, 'out', 'presentations', DATE_TAG, COUNTRY)
OUT_FILE   = os.path.join(OUT_DIR, f'{COUNTRY}_labor_trends_{DATE_TAG}.pptx')

# Slide geometry (cm) — widescreen 16:9
SLIDE_W  = 33.87
SLIDE_H  = 19.05

# Content area boundaries
TITLE_BOTTOM = 1.9      # cm — bottom of title box
FOOTER_TOP   = SLIDE_H - 0.8
AVAIL_H      = FOOTER_TOP - TITLE_BOTTOM
AVAIL_W      = SLIDE_W - 2.0

# Default chart width (will be scaled down if needed)
DEFAULT_IMG_W = 22.0

# Colors
C_BLUE    = RGBColor(31,  73,  125)   # IDB Blue  #1F497D
C_LBLUE   = 'DCE6F1'                  # IDB Light Blue (hex for lxml)
C_GREY    = RGBColor(89,  89,  89)
C_LGREY   = RGBColor(120, 120, 120)
C_WHITE   = RGBColor(255, 255, 255)
C_BLACK   = RGBColor(0,   0,   0)

FOOTER_TEXT = 'HDMF — How do Migrants Fare in LAC | IDB'

# Chart order: all COL charts land directly in TREND_DIR (no labor/ subfolder)
CHART_ORDER = [
    ('labor_employment_trend.png',          'Tasa de empleo (16-64)',               'trend'),
    ('labor_unemployment_trend.png',        'Tasa de desempleo (16-64)',             'trend'),
    ('labor_inactivity_trend.png',          'Tasa de inactividad (16-64)',           'trend'),
    ('labor_pea_trend.png',                 'Tasa PEA (16-64)',                      'trend'),
    ('formality_trend.png',                 'Tasa de formalidad',                    'trend'),
    ('wages_income_trend.png',              'Ingreso total promedio',                'trend'),
    ('wages_hours_trend.png',               'Horas totales trabajadas (promedio)',   'trend'),
    ('population_migrant_share_trend.png',  'Participación de migrantes',            'trend'),
]

# Timing table: header + data rows
# Columns: #, Etapa, Proceso, Descripción breve, Resultado, Tiempo
TIMING_ROWS = [
    ('#', 'Etapa', 'Proceso', 'Descripción', 'Resultado principal', 'Tiempo'),
    ('1', 'Armonización\n(Stata)',
          'COL_YYYY_variablesBID.do\n(ejecución en paralelo)',
          'Traduce variables de la encuesta cruda al esquema HDMF. Se ejecuta una ola por script; varias olas en paralelo.',
          '*_BID.dta por ola',
          '~3 min\n(todas las olas)'),
    ('2', 'Caché\n(Python)',
          'hdmf_build.py',
          'Carga todos los *_BID.dta (~13 GB), estandariza tipos, agrega periodo_c y guarda un pickle unificado.',
          '_hdmf_cache.pkl',
          '~10-15 min\n(solo si hay datos nuevos)'),
    ('3', 'Indicadores\nde tendencia\n(Python)',
          'hdmf_trends.py\n--country COL',
          'Lee el caché, filtra por país y período, calcula indicadores ponderados y exporta a Excel.',
          'hdmf_trends_COL.xlsx\n(11 hojas)',
          '~30-60 s'),
    ('4', 'Gráficos\nde tendencia\n(Python)',
          'hdmf_trends_charts.py\n--country COL',
          'Genera un gráfico PNG por indicador con los períodos en el eje X y grupo migratorio como hue.',
          '~25 gráficos PNG\nen out/trends/',
          '~15 s'),
]

# =============================================================================
# HELPERS
# =============================================================================

def _set_font(run, size_pt, bold=False, color=None):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = 'Calibri'
    if color:
        run.font.color.rgb = color


def _cell_fill(cell, hex_color):
    """Set solid fill on a table cell using lxml."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    sf   = etree.SubElement(tcPr, qn('a:solidFill'))
    sc   = etree.SubElement(sf,   qn('a:srgbClr'))
    sc.set('val', hex_color)


def add_footer(slide):
    txb = slide.shapes.add_textbox(Cm(1), Cm(FOOTER_TOP), Cm(SLIDE_W - 2), Cm(0.6))
    tf  = txb.text_frame
    p   = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = FOOTER_TEXT
    _set_font(run, 8, color=C_LGREY)


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_title_box(slide, text):
    txb = slide.shapes.add_textbox(Cm(1), Cm(0.4), Cm(SLIDE_W - 2), Cm(1.4))
    tf  = txb.text_frame
    p   = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_font(run, 18, bold=True, color=C_BLUE)


# =============================================================================
# SLIDE BUILDERS
# =============================================================================

def add_title_slide(prs):
    slide = _blank_slide(prs)

    txb = slide.shapes.add_textbox(Cm(3), Cm(5), Cm(28), Cm(3))
    tf  = txb.text_frame
    p   = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'Mercado Laboral — Tendencias'
    _set_font(run, 32, bold=True, color=C_BLUE)

    txb2 = slide.shapes.add_textbox(Cm(3), Cm(8.5), Cm(28), Cm(2))
    p2   = txb2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = f'{COUNTRY}  |  2018 - 2025'
    _set_font(run2, 20, color=C_GREY)

    txb3 = slide.shapes.add_textbox(Cm(3), Cm(11), Cm(28), Cm(1.5))
    p3   = txb3.text_frame.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    run3 = p3.add_run()
    run3.text = f'HDMF — How do Migrants Fare in LAC  |  IDB  |  {DATE_TAG}'
    _set_font(run3, 12, color=C_LGREY)


def add_chart_slide(prs, img_path, title):
    """One chart per slide, image centered both horizontally and vertically."""
    slide = _blank_slide(prs)
    _add_title_box(slide, title)

    with PILImage.open(img_path) as im:
        px_w, px_h = im.size
    ratio = px_h / px_w

    img_w = DEFAULT_IMG_W
    img_h = img_w * ratio

    # Scale to fit available area
    if img_h > AVAIL_H:
        img_h = AVAIL_H
        img_w = img_h / ratio
    if img_w > AVAIL_W:
        img_w = AVAIL_W
        img_h = img_w * ratio

    # Center: horizontal and vertical
    img_left = (SLIDE_W - img_w) / 2
    img_top  = TITLE_BOTTOM + (AVAIL_H - img_h) / 2

    slide.shapes.add_picture(img_path, Cm(img_left), Cm(img_top), Cm(img_w), Cm(img_h))
    add_footer(slide)


def _fmt_pct(v):
    """Format a float as a percentage string (e.g. 0.752 -> '75.2%')."""
    try:
        return f'{float(v)*100:.1f}%'
    except Exception:
        return str(v)


def add_labor_table_slides(prs, excel_path):
    """Two slides with labor rate tables — clean, no fill colors."""
    df = pd.read_excel(excel_path, sheet_name='labor_rate', index_col=[0, 1])
    periods = list(df.columns)

    # Short period labels for column headers
    period_labels = [p.replace('t3', ' T3').replace('m', '-M') for p in periods]

    groups = ['Nativos', 'Migrantes de Venezuela', 'Migrantes de otros países']

    slide_specs = [
        {
            'title': 'Indicadores laborales (I) — Empleo y Desempleo',
            'variables': ['Tasa de empleo', 'Tasa de desempleo'],
        },
        {
            'title': 'Indicadores laborales (II) — Inactividad y PEA',
            'variables': ['Tasa de inactividad', 'Tasa PEA'],
        },
    ]

    for spec in slide_specs:
        slide = _blank_slide(prs)
        _add_title_box(slide, spec['title'])
        add_footer(slide)

        header = ['Grupo', 'Indicador'] + period_labels
        rows   = [header]

        for var in spec['variables']:
            for grp in groups:
                try:
                    vals = df.loc[(grp, var)]
                    row_vals = [_fmt_pct(vals.get(p, '')) for p in periods]
                except KeyError:
                    row_vals = [''] * len(periods)
                rows.append([grp, var] + row_vals)

        n_rows = len(rows)
        n_cols = len(header)

        tbl_left  = 1.0
        tbl_top   = 2.2
        tbl_w     = SLIDE_W - 2.0
        tbl_h     = FOOTER_TOP - tbl_top - 0.3

        table = slide.shapes.add_table(
            n_rows, n_cols, Cm(tbl_left), Cm(tbl_top), Cm(tbl_w), Cm(tbl_h)
        ).table

        remaining   = 1.0 - 0.14 - 0.18
        period_frac = remaining / len(periods)
        col_fracs   = [0.14, 0.18] + [period_frac] * len(periods)
        for ci, frac in enumerate(col_fracs):
            table.columns[ci].width = Cm(tbl_w * frac)

        for ri, row_data in enumerate(rows):
            is_header = (ri == 0)
            for ci, cell_text in enumerate(row_data):
                cell = table.cell(ri, ci)
                cell.text = str(cell_text)
                tf2 = cell.text_frame
                tf2.word_wrap = True
                p2  = tf2.paragraphs[0]
                is_num = ci >= 2 and not is_header
                p2.alignment = PP_ALIGN.RIGHT if is_num else PP_ALIGN.LEFT
                run2 = p2.runs[0] if p2.runs else p2.add_run()
                run2.font.name = 'Calibri'
                run2.font.size = Pt(9 if not is_header else 10)
                run2.font.bold = is_header
                run2.font.color.rgb = C_BLUE if is_header else C_BLACK

        print(f'  Added: labor table — {spec["title"]}')


def add_timing_table_slide(prs):
    """Pipeline timing summary with stage descriptions."""
    slide = _blank_slide(prs)
    _add_title_box(slide, 'Pipeline HDMF — Etapas y tiempos de proceso')
    add_footer(slide)

    rows   = TIMING_ROWS
    n_rows = len(rows)
    n_cols = len(rows[0])

    tbl_left = 1.0
    tbl_top  = 2.2
    tbl_w    = SLIDE_W - 2.0
    tbl_h    = FOOTER_TOP - tbl_top - 0.3

    table = slide.shapes.add_table(
        n_rows, n_cols, Cm(tbl_left), Cm(tbl_top), Cm(tbl_w), Cm(tbl_h)
    ).table

    # Column fractions: #, Etapa, Proceso, Descripcion, Resultado, Tiempo
    col_fracs = [0.04, 0.10, 0.14, 0.38, 0.20, 0.14]
    for ci, frac in enumerate(col_fracs):
        table.columns[ci].width = Cm(tbl_w * frac)

    for ri, row_data in enumerate(rows):
        is_header = (ri == 0)
        for ci, cell_text in enumerate(row_data):
            cell = table.cell(ri, ci)
            cell.text = str(cell_text)
            tf2 = cell.text_frame
            tf2.word_wrap = True
            p2  = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.LEFT
            run2 = p2.runs[0] if p2.runs else p2.add_run()
            run2.font.name = 'Calibri'
            run2.font.size = Pt(9 if not is_header else 10)
            run2.font.bold = is_header
            run2.font.color.rgb = C_BLUE if is_header else C_BLACK

    print('  Added: pipeline timing table')


# =============================================================================
# MAIN
# =============================================================================

def main():
    os.makedirs(OUT_DIR, exist_ok=True)

    # Resolve chart paths
    charts = []
    for fname, title, folder in CHART_ORDER:
        base = LABOR_DIR if folder == 'labor' else TREND_DIR
        path = os.path.join(base, fname)
        if os.path.exists(path):
            charts.append((path, title))
        else:
            print(f'  [SKIP] Not found: {fname}')

    print(f'Charts found : {len(charts)}')

    prs = Presentation()
    prs.slide_width  = Cm(SLIDE_W)
    prs.slide_height = Cm(SLIDE_H)

    # 1 — Title
    add_title_slide(prs)

    # 2 — One chart per slide
    for img_path, title in charts:
        add_chart_slide(prs, img_path, title)
        print(f'  Added: {os.path.basename(img_path)}')

    # 3 — Labor summary tables (2 slides)
    if os.path.exists(EXCEL_FILE):
        add_labor_table_slides(prs, EXCEL_FILE)
    else:
        print(f'  [SKIP] Excel not found: {EXCEL_FILE}')

    # 4 — Pipeline timing
    add_timing_table_slide(prs)

    prs.save(OUT_FILE)
    print(f'\nSaved -> {OUT_FILE}')
    print(f'Total slides : {len(prs.slides)}')


if __name__ == '__main__':
    main()
