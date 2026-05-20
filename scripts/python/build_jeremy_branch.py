"""
build_jeremy_branch.py  —  Jeremy's branch outputs
====================================================
Output: out/jeremy/2026-05-11/jeremy_indicators.xlsx

Sheet layout:
  unemployment    One table per country (stacked). Rows = year, cols = Native | Migrant
  informality     same
  inactivity      same
  wap             same
  education_abs   One table per country (stacked). Rows = edu category, cols = Native | Migrant  (weighted counts)
  education_pct   same but % share

Countries: COL, CHL, ECU, PER
Groups: Native / Migrant  (all migrants combined)
Education: most recent wave per country only.
"""

import os, pickle, numpy as np, pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Paths & config
# ---------------------------------------------------------------------------
BASE = Path(__file__).parent.parent
CACHE_FILE   = BASE / "bases armo" / "armo" / "_hdmf_cache.pkl"
DATE_TAG     = "2026-05-11"
OUT_DIR      = BASE / "out" / "jeremy" / DATE_TAG
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_FILE     = OUT_DIR / "jeremy_indicators.xlsx"

LAC4         = ["COL", "CHL", "ECU", "PER"]
RECENT_WAVES = {"COL": "2025t3", "CHL": "2024a", "ECU": "2025m12", "PER": "2024a"}
GROUPS       = ["Native", "Migrant"]
GROUP_COL    = "migrante_ci"

# ---------------------------------------------------------------------------
# Load cache
# ---------------------------------------------------------------------------
print("Loading cache …")
with open(CACHE_FILE, "rb") as fh:
    _cache = pickle.load(fh)
data = _cache["data"]

df = data[data["pais_c"].isin(LAC4)].copy()
df["year"] = df["periodo_c"].str.extract(r"(\d{4})").astype(int)

# ---------------------------------------------------------------------------
# Derived indicators
# ---------------------------------------------------------------------------
wa_mask = df["edad_ci"].between(16, 64)

df["wap_ci"]   = df["edad_ci"].between(15, 64).astype(int)
df["inact_ci"] = df["inactivo_ci"].where(wa_mask, np.nan)
df["desemp_r"] = (df["desemp_ci"]
                  .where(wa_mask, np.nan)
                  .where(df["pea_ci"] == 1, np.nan))
df["informal"] = np.where(df["emp_ci"] == 1,
                          1 - pd.to_numeric(df["formal_ci"], errors="coerce"),
                          np.nan)

INDICATORS = {
    "unemployment": ("desemp_r",  "Unemployment rate (% of active pop.)"),
    "informality":  ("informal",  "Informality rate (% of employed)"),
    "inactivity":   ("inact_ci",  "Inactivity rate (% of WAP 16–64)"),
    "wap":          ("wap_ci",    "Working-age pop. (% of total pop., 15–64)"),
}

# ---------------------------------------------------------------------------
# Weighted mean helper
# ---------------------------------------------------------------------------
def _wavg(series, weights):
    v = pd.to_numeric(series, errors="coerce").to_numpy(float)
    w = pd.to_numeric(weights, errors="coerce").to_numpy(float)
    mask = np.isfinite(v) & np.isfinite(w) & (w > 0)
    if mask.sum() == 0:
        return np.nan
    v_, w_ = v[mask], w[mask]
    return float(np.dot(w_, v_) / w_.sum())

# ---------------------------------------------------------------------------
# 1. Build trend data  (year × group per country per indicator)
# ---------------------------------------------------------------------------
print("Computing trend indicators …")
trend = {}   # trend[ind_key][country] = DataFrame(index=period, cols=[Native, Migrant])

for ind_key, (col, _) in INDICATORS.items():
    trend[ind_key] = {}
    for country in LAC4:
        dfc = df[df["pais_c"] == country]
        rows = []
        for period in sorted(dfc["periodo_c"].unique()):
            yr  = int(period[:4])
            dfp = dfc[dfc["periodo_c"] == period]
            row = {"Period": period, "Year": yr}
            for grp in GROUPS:
                mu = _wavg(dfp.loc[dfp[GROUP_COL] == grp, col],
                           dfp.loc[dfp[GROUP_COL] == grp, "factor_ci"])
                row[grp] = round(mu * 100, 2) if pd.notna(mu) else np.nan
            rows.append(row)
        trend[ind_key][country] = pd.DataFrame(rows).set_index(["Period", "Year"])

# ---------------------------------------------------------------------------
# 2. Build education data  (category × group per country, most recent wave)
# ---------------------------------------------------------------------------
EDU_3CAT_MAP = {
    1: "Primary or less", 2: "Primary or less", 3: "Primary or less",
    4: "Secondary",       5: "Secondary",       6: "Secondary",
    7: "Higher education", 8: "Higher education",
    9: "Higher education", 10: "Higher education",
}
EDU_ORDER = ["Primary or less", "Secondary", "Higher education", "Total"]

print("Computing education distribution …")
edu_abs = {}   # edu_abs[country] = DataFrame
edu_pct = {}

for country, period in RECENT_WAVES.items():
    dfc = df[(df["pais_c"] == country) & (df["periodo_c"] == period)].copy()
    dfc["edu_3cat"] = pd.to_numeric(dfc["edu_hdmf"], errors="coerce").map(EDU_3CAT_MAP)
    dfc = dfc[dfc["edu_3cat"].notna() & dfc["factor_ci"].notna()]

    abs_rows, pct_rows = [], []
    for cat in EDU_ORDER[:-1]:
        abs_row = {"Education level": cat}
        pct_row = {"Education level": cat}
        for grp in GROUPS:
            dfg    = dfc[dfc[GROUP_COL] == grp]
            total  = dfg["factor_ci"].sum()
            cat_w  = dfg.loc[dfg["edu_3cat"] == cat, "factor_ci"].sum()
            abs_row[grp] = int(round(cat_w))
            pct_row[grp] = round(cat_w / total * 100, 2) if total > 0 else np.nan
        abs_rows.append(abs_row)
        pct_rows.append(pct_row)

    # Total row
    abs_tot = {"Education level": "Total"}
    pct_tot = {"Education level": "Total"}
    for grp in GROUPS:
        dfg = dfc[dfc[GROUP_COL] == grp]
        tot = dfg["factor_ci"].sum()
        abs_tot[grp] = int(round(tot))
        pct_tot[grp] = 100.0
    abs_rows.append(abs_tot)
    pct_rows.append(pct_tot)

    edu_abs[country] = pd.DataFrame(abs_rows).set_index("Education level")
    edu_pct[country] = pd.DataFrame(pct_rows).set_index("Education level")

# ---------------------------------------------------------------------------
# 3. Write to Excel using openpyxl directly
# ---------------------------------------------------------------------------
# Styles
C_NAVY    = "1D3557"
C_BLUE    = "2E75B6"
C_ORANGE  = "C55A11"
C_GREEN   = "375623"
C_PURPLE  = "7030A0"
C_YELLOW  = "7F6000"

IND_COLORS = {
    "unemployment": ("FDE9D9", C_ORANGE),   # light orange / dark orange header
    "informality":  ("E2EFDA", C_GREEN),
    "inactivity":   ("DEEAF1", C_BLUE),
    "wap":          ("F2E7FA", C_PURPLE),
}
EDU_COLOR    = ("FFF2CC", C_YELLOW)
COUNTRY_FILLS = {
    "COL": "EBF5FB",
    "CHL": "EBF5EB",
    "ECU": "FFFBEC",
    "PER": "FBEBED",
}

def _fill(hex_col):
    return PatternFill("solid", fgColor=hex_col)

def _thin():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _med_bottom():
    med = Side(style="medium", color="888888")
    thn = Side(style="thin",   color="CCCCCC")
    return Border(left=thn, right=thn, top=thn, bottom=med)

def write_country_table(ws, start_row, country, period, df_table,
                        row_label_col, header_fill, country_fill,
                        num_format="0.00", title=None):
    """
    Write a single country sub-table starting at start_row.
    Returns the next free row after the table.
    """
    cols = df_table.columns.tolist()
    n_cols = len(cols) + 1   # +1 for the row-label column

    # Country title row
    title_text = title or f"{country}  ({period})"
    ws.cell(start_row, 1, title_text).font      = Font(bold=True, size=11, color="FFFFFF")
    ws.cell(start_row, 1).fill                  = _fill(C_NAVY)
    ws.cell(start_row, 1).alignment             = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[start_row].height         = 20
    for c in range(2, n_cols + 1):
        ws.cell(start_row, c).fill = _fill(C_NAVY)

    # Header row
    hr = start_row + 1
    ws.cell(hr, 1, row_label_col).font      = Font(bold=True, size=10, color="FFFFFF")
    ws.cell(hr, 1).fill                     = _fill(header_fill)
    ws.cell(hr, 1).alignment                = Alignment(horizontal="left", vertical="center")
    ws.cell(hr, 1).border                   = _thin()
    ws.row_dimensions[hr].height            = 18
    for ci, col_name in enumerate(cols, start=2):
        cell = ws.cell(hr, ci, col_name)
        cell.font      = Font(bold=True, size=10, color="FFFFFF")
        cell.fill      = _fill(header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _thin()

    # Data rows
    for ri, (idx, row) in enumerate(df_table.iterrows()):
        dr    = hr + 1 + ri
        shade = country_fill if ri % 2 == 0 else "FFFFFF"
        is_total = str(idx) == "Total" if not isinstance(idx, tuple) else False

        # Row label
        lbl_cell = ws.cell(dr, 1)
        lbl_val  = " / ".join(str(i) for i in idx) if isinstance(idx, tuple) else str(idx)
        lbl_cell.value     = lbl_val
        lbl_cell.font      = Font(bold=is_total, size=9)
        lbl_cell.fill      = _fill("F0F0F0") if is_total else _fill(shade)
        lbl_cell.alignment = Alignment(horizontal="left", vertical="center")
        lbl_cell.border    = _med_bottom() if is_total else _thin()

        for ci, val in enumerate(row, start=2):
            dc = ws.cell(dr, ci, val if pd.notna(val) else None)
            dc.font          = Font(bold=is_total, size=9)
            dc.fill          = _fill("F0F0F0") if is_total else _fill(shade)
            dc.alignment     = Alignment(horizontal="right", vertical="center")
            dc.number_format = num_format
            dc.border        = _med_bottom() if is_total else _thin()
        ws.row_dimensions[dr].height = 14

    next_row = hr + 1 + len(df_table) + 2   # 2 blank rows gap
    return next_row

wb = Workbook()
wb.remove(wb.active)   # remove default sheet

# ── Indicator sheets ─────────────────────────────────────────────────────────
for ind_key, (_, label) in INDICATORS.items():
    bg_light, hdr_color = IND_COLORS[ind_key]
    ws = wb.create_sheet(title=ind_key)

    # Sheet title (row 1)
    ws.cell(1, 1, label).font      = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(1, 1).fill             = _fill(C_NAVY)
    ws.cell(1, 1).alignment        = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height    = 24
    for c in range(2, 5):
        ws.cell(1, c).fill = _fill(C_NAVY)

    current_row = 3
    for country in LAC4:
        period    = RECENT_WAVES[country]
        tbl       = trend[ind_key][country]
        c_fill    = COUNTRY_FILLS[country]
        current_row = write_country_table(
            ws, current_row, country, f"all waves — latest: {period}",
            tbl, "Period / Year", hdr_color, c_fill,
            num_format="0.00",
        )

    # Column widths
    ws.column_dimensions["A"].width = 16
    for c in ["B", "C"]:
        ws.column_dimensions[c].width = 14

# ── Education sheets ──────────────────────────────────────────────────────────
for sheet_label, data_dict, num_fmt, sheet_title in [
    ("education_abs", edu_abs, "#,##0",  "Education level — weighted population count"),
    ("education_pct", edu_pct, "0.00",   "Education level — % share within group"),
]:
    bg_light, hdr_color = EDU_COLOR
    ws = wb.create_sheet(title=sheet_label)

    ws.cell(1, 1, sheet_title).font   = Font(bold=True, size=12, color="FFFFFF")
    ws.cell(1, 1).fill                = _fill(C_NAVY)
    ws.cell(1, 1).alignment           = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height       = 24
    for c in range(2, 5):
        ws.cell(1, c).fill = _fill(C_NAVY)

    current_row = 3
    for country in LAC4:
        period    = RECENT_WAVES[country]
        tbl       = data_dict[country]
        c_fill    = COUNTRY_FILLS[country]
        current_row = write_country_table(
            ws, current_row, country, period,
            tbl, "Education level", hdr_color, c_fill,
            num_format=num_fmt,
        )

    ws.column_dimensions["A"].width = 22
    for c in ["B", "C"]:
        ws.column_dimensions[c].width = 18

wb.save(OUT_FILE)

# ---------------------------------------------------------------------------
# Summary
# ---------------------------------------------------------------------------
print(f"\nDone -> {OUT_FILE}")
print(f"  Size: {os.path.getsize(OUT_FILE)/1e3:.0f} KB")
print(f"\nSheets: {[s.title for s in wb.worksheets]}")

print("\nEducation % (most recent wave):")
for country, period in RECENT_WAVES.items():
    print(f"\n  {country} ({period})")
    print(edu_pct[country].to_string())
